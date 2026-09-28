"""
reports.py — Excel report generation (per class, with attendance %)
"""
import logging
from datetime import datetime
from pathlib import Path

from app.config import REPORTS_DIR
from app.db import run_query

logger = logging.getLogger(__name__)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


def _border():
    s = Side(border_style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)


def generate_class_report(class_id: int, cls_data: dict) -> str:
    """
    Generate Excel attendance report for a class.
    Returns the file path (str).
    """
    if not HAS_OPENPYXL:
        raise RuntimeError("openpyxl not installed")

    sessions = run_query(
        "SELECT * FROM sessions WHERE class_id=%s ORDER BY started_at;", (class_id,)
    )
    students = run_query("""
        SELECT p.id, p.name, p.roll_no
        FROM class_students cs
        JOIN people p ON p.id = cs.person_id
        WHERE cs.class_id=%s ORDER BY p.name;
    """, (class_id,))

    total = len(sessions)
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    # ── Title row ──────────────────────────────────────────────
    title_fill = PatternFill("solid", fgColor="1E3A5F")
    ws.merge_cells("A1:Z1")
    cell = ws["A1"]
    cell.value = f"{cls_data['name']} ({cls_data['code']}) — Attendance Report"
    cell.font = Font(bold=True, color="FFFFFF", size=13)
    cell.fill = title_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # ── Sub-header ─────────────────────────────────────────────
    ws.merge_cells("A2:Z2")
    ws["A2"].value = f"Generated: {datetime.now().strftime('%d %b %Y %I:%M %p')} | Total Sessions: {total}"
    ws["A2"].font = Font(italic=True, color="555555")
    ws["A2"].alignment = Alignment(horizontal="center")

    # ── Column headers ─────────────────────────────────────────
    headers = ["#", "Roll No", "Student Name"]
    for s in sessions:
        dt = s["started_at"]
        if isinstance(dt, str):
            dt = datetime.fromisoformat(dt.replace("Z", ""))
        headers.append(dt.strftime("%d/%m"))
    headers += ["Present", "Total", "Attendance %", "Status"]

    header_fill = PatternFill("solid", fgColor="2563EB")
    row = 3
    for col_idx, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col_idx, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = _border()

    # ── Data rows ──────────────────────────────────────────────
    for i, student in enumerate(students, 1):
        row += 1
        present_sessions = {
            r["session_id"]
            for r in run_query(
                "SELECT session_id FROM attendance_log WHERE person_id=%s AND class_id=%s;",
                (student["id"], class_id),
            )
        }
        attended = len(present_sessions)
        pct = round(attended / total * 100, 1) if total > 0 else 0.0
        status = "Good" if pct >= 75 else ("Low" if pct >= 50 else "Critical")

        fill_color = (
            "D1FAE5" if pct >= 75 else ("FEF3C7" if pct >= 50 else "FEE2E2")
        )
        row_fill = PatternFill("solid", fgColor=fill_color)

        data = [i, student.get("roll_no") or "—", student["name"]]
        for s in sessions:
            data.append("P" if s["id"] in present_sessions else "A")
        data += [attended, total, f"{pct}%", status]

        for col_idx, val in enumerate(data, 1):
            c = ws.cell(row=row, column=col_idx, value=val)
            c.alignment = Alignment(horizontal="center")
            c.border = _border()
            if col_idx in (len(data) - 1, len(data)):  # % and Status cols
                c.fill = row_fill

    # ── Column widths ──────────────────────────────────────────
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 28
    for idx in range(4, len(sessions) + 4):
        ws.column_dimensions[get_column_letter(idx)].width = 7
    ws.column_dimensions[get_column_letter(len(sessions) + 4)].width = 9
    ws.column_dimensions[get_column_letter(len(sessions) + 5)].width = 7
    ws.column_dimensions[get_column_letter(len(sessions) + 6)].width = 13
    ws.column_dimensions[get_column_letter(len(sessions) + 7)].width = 11

    fname = f"attendance_{cls_data['code']}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    fpath = REPORTS_DIR / fname
    wb.save(str(fpath))
    logger.info(f"Report saved: {fpath}")
    return str(fpath)
